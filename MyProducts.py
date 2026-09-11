import sqlite3
from contextlib import contextmanager
from pathlib import Path

from openpyxl import Workbook
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QApplication,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)


class ProductDatabase:
    def __init__(self, database_path="products.db"):
        self.database_path = Path(database_path)
        self.create_table()

    @contextmanager
    def connect(self):
        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        try:
            yield connection
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def create_table(self):
        with self.connect() as connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS Products (
                    productID INTEGER PRIMARY KEY AUTOINCREMENT,
                    productName TEXT NOT NULL,
                    productPrice INTEGER NOT NULL CHECK (productPrice >= 0)
                )
                """
            )

    def add_product(self, product_name, product_price):
        with self.connect() as connection:
            cursor = connection.execute(
                """
                INSERT INTO Products (productName, productPrice)
                VALUES (?, ?)
                """,
                (product_name, product_price),
            )
            return cursor.lastrowid

    def update_product(self, product_id, product_name, product_price):
        with self.connect() as connection:
            cursor = connection.execute(
                """
                UPDATE Products
                SET productName = ?, productPrice = ?
                WHERE productID = ?
                """,
                (product_name, product_price, product_id),
            )
            return cursor.rowcount > 0

    def delete_product(self, product_id):
        with self.connect() as connection:
            cursor = connection.execute(
                "DELETE FROM Products WHERE productID = ?",
                (product_id,),
            )
            return cursor.rowcount > 0

    def delete_products(self, product_ids):
        if not product_ids:
            return 0

        placeholders = ", ".join("?" for _ in product_ids)
        with self.connect() as connection:
            cursor = connection.execute(
                f"DELETE FROM Products WHERE productID IN ({placeholders})",
                tuple(product_ids),
            )
            return cursor.rowcount

    def search_products(self, keyword=""):
        with self.connect() as connection:
            if keyword:
                cursor = connection.execute(
                    """
                    SELECT productID, productName, productPrice
                    FROM Products
                    WHERE productName LIKE ?
                    ORDER BY productID
                    """,
                    (f"%{keyword}%",),
                )
            else:
                cursor = connection.execute(
                    """
                    SELECT productID, productName, productPrice
                    FROM Products
                    ORDER BY productID
                    """
                )
            return [dict(row) for row in cursor.fetchall()]


class ProductWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.database = ProductDatabase()
        self.products = []
        self.setWindowTitle("ERP 상품 관리")
        self.resize(900, 620)
        self.setup_ui()
        self.load_products()

    def setup_ui(self):
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)
        main_layout.setContentsMargins(22, 18, 22, 18)
        main_layout.setSpacing(14)

        self.setStyleSheet(
            """
            QMainWindow, QWidget {
                background-color: #f3f5f7;
                color: #263238;
                font-family: 'Malgun Gothic';
                font-size: 10pt;
            }
            QLabel#pageTitle {
                color: #17365d;
                font-size: 19pt;
                font-weight: 700;
            }
            QLabel#pageSubtitle, QLabel#sectionTitle {
                color: #64748b;
            }
            QLabel#sectionTitle {
                color: #17365d;
                font-weight: 700;
            }
            QFrame#panel {
                background-color: #ffffff;
                border: 1px solid #d9e0e7;
                border-radius: 4px;
            }
            QLineEdit, QSpinBox {
                background-color: #ffffff;
                border: 1px solid #bcc8d4;
                border-radius: 3px;
                padding: 7px 8px;
                min-height: 18px;
            }
            QLineEdit:focus, QSpinBox:focus {
                border: 1px solid #2878b5;
            }
            QPushButton {
                background-color: #ffffff;
                border: 1px solid #aebdca;
                border-radius: 3px;
                color: #263238;
                padding: 7px 13px;
                min-height: 18px;
            }
            QPushButton:hover {
                background-color: #eaf2f8;
                border-color: #2878b5;
            }
            QPushButton#primaryButton {
                background-color: #1f5f8b;
                border-color: #1f5f8b;
                color: #ffffff;
                font-weight: 700;
            }
            QPushButton#dangerButton {
                color: #b42318;
                border-color: #d9a6a1;
            }
            QTableWidget {
                background-color: #ffffff;
                alternate-background-color: #f7f9fb;
                border: 1px solid #d9e0e7;
                gridline-color: #e4e9ee;
                selection-background-color: #dbeaf5;
                selection-color: #263238;
            }
            QHeaderView::section {
                background-color: #e7edf3;
                border: 0;
                border-bottom: 1px solid #cbd5df;
                color: #34495e;
                font-weight: 700;
                padding: 8px;
            }
            """
        )

        title = QLabel("상품 관리")
        title.setObjectName("pageTitle")
        subtitle = QLabel("기준정보 관리  >  상품 마스터")
        subtitle.setObjectName("pageSubtitle")
        main_layout.addWidget(title)
        main_layout.addWidget(subtitle)

        input_panel = QWidget()
        input_panel.setObjectName("panel")
        input_layout = QVBoxLayout(input_panel)
        input_layout.setContentsMargins(16, 14, 16, 14)
        section_title = QLabel("상품 정보 입력")
        section_title.setObjectName("sectionTitle")
        input_layout.addWidget(section_title)
        form_layout = QFormLayout()
        form_layout.setContentsMargins(0, 4, 0, 0)
        form_layout.setHorizontalSpacing(18)
        self.name_input = QLineEdit()
        self.name_input.setPlaceholderText("상품명을 입력하세요")
        self.price_input = QSpinBox()
        self.price_input.setRange(0, 2_147_483_647)
        self.price_input.setSuffix(" 원")
        form_layout.addRow("상품명", self.name_input)
        form_layout.addRow("상품가격", self.price_input)
        input_layout.addLayout(form_layout)

        button_layout = QHBoxLayout()
        self.add_button = QPushButton("상품 입력")
        self.update_button = QPushButton("상품 수정")
        self.delete_button = QPushButton("선택 삭제")
        self.clear_button = QPushButton("입력 지우기")
        self.add_button.setObjectName("primaryButton")
        self.delete_button.setObjectName("dangerButton")
        button_layout.addWidget(self.add_button)
        button_layout.addWidget(self.update_button)
        button_layout.addWidget(self.delete_button)
        button_layout.addWidget(self.clear_button)
        button_layout.addStretch()
        input_layout.addLayout(button_layout)
        main_layout.addWidget(input_panel)

        list_panel = QWidget()
        list_panel.setObjectName("panel")
        list_layout = QVBoxLayout(list_panel)
        list_layout.setContentsMargins(16, 14, 16, 14)
        list_header_layout = QHBoxLayout()
        list_title = QLabel("상품 목록")
        list_title.setObjectName("sectionTitle")
        list_header_layout.addWidget(list_title)
        list_header_layout.addStretch()
        list_header_layout.addWidget(QLabel("체크한 상품을 일괄 삭제할 수 있습니다."))
        list_layout.addLayout(list_header_layout)

        search_layout = QHBoxLayout()
        search_layout.addWidget(QLabel("상품명"))
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("상품명을 입력하세요. 비워두면 전체 검색")
        self.search_button = QPushButton("검색")
        self.excel_button = QPushButton("엑셀로 저장")
        search_layout.addWidget(self.search_input)
        search_layout.addWidget(self.search_button)
        search_layout.addWidget(self.excel_button)
        list_layout.addLayout(search_layout)

        self.product_table = QTableWidget(0, 4)
        self.product_table.setHorizontalHeaderLabels(["선택", "상품 ID", "상품명", "상품가격"])
        self.product_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.product_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.product_table.setAlternatingRowColors(True)
        self.product_table.setColumnWidth(0, 55)
        self.product_table.setColumnWidth(1, 100)
        self.product_table.setColumnWidth(2, 320)
        self.product_table.horizontalHeader().setStretchLastSection(True)
        list_layout.addWidget(self.product_table)
        self.status_label = QLabel("총 0건")
        self.status_label.setObjectName("pageSubtitle")
        list_layout.addWidget(self.status_label)
        main_layout.addWidget(list_panel, 1)

        self.add_button.clicked.connect(self.add_product)
        self.update_button.clicked.connect(self.update_product)
        self.delete_button.clicked.connect(self.delete_product)
        self.clear_button.clicked.connect(self.clear_inputs)
        self.search_button.clicked.connect(self.load_products)
        self.search_input.returnPressed.connect(self.load_products)
        self.excel_button.clicked.connect(self.export_to_excel)
        self.product_table.itemSelectionChanged.connect(self.select_product)

    def load_products(self):
        keyword = self.search_input.text().strip()
        self.products = self.database.search_products(keyword)
        self.product_table.setRowCount(len(self.products))
        self.status_label.setText(f"총 {len(self.products):,}건")

        for row_index, product in enumerate(self.products):
            values = (
                "",
                product["productID"],
                product["productName"],
                f"{product['productPrice']:,}원",
            )
            for column_index, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                if column_index == 0:
                    item.setFlags(
                        item.flags() | Qt.ItemFlag.ItemIsUserCheckable
                    )
                    item.setCheckState(Qt.CheckState.Unchecked)
                    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                if column_index in (1, 3):
                    item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.product_table.setItem(row_index, column_index, item)

    def checked_product_ids(self):
        product_ids = []
        for row_index, product in enumerate(self.products):
            checkbox = self.product_table.item(row_index, 0)
            if checkbox and checkbox.checkState() == Qt.CheckState.Checked:
                product_ids.append(product["productID"])
        return product_ids

    def selected_product_id(self):
        selected_rows = self.product_table.selectionModel().selectedRows()
        if not selected_rows:
            return None
        return self.products[selected_rows[0].row()]["productID"]

    def select_product(self):
        selected_rows = self.product_table.selectionModel().selectedRows()
        if not selected_rows:
            return
        product = self.products[selected_rows[0].row()]
        self.name_input.setText(product["productName"])
        self.price_input.setValue(product["productPrice"])

    def add_product(self):
        name = self.name_input.text().strip()
        if not name:
            QMessageBox.warning(self, "입력 오류", "상품명을 입력하세요.")
            return

        self.database.add_product(name, self.price_input.value())
        self.clear_inputs()
        self.load_products()
        self.name_input.setFocus()

    def update_product(self):
        product_id = self.selected_product_id()
        name = self.name_input.text().strip()
        if product_id is None:
            QMessageBox.warning(self, "선택 오류", "수정할 상품을 목록에서 선택하세요.")
            return
        if not name:
            QMessageBox.warning(self, "입력 오류", "상품명을 입력하세요.")
            return

        self.database.update_product(product_id, name, self.price_input.value())
        self.clear_inputs()
        self.load_products()

    def delete_product(self):
        product_ids = self.checked_product_ids()
        if not product_ids:
            QMessageBox.warning(self, "선택 오류", "삭제할 상품을 체크하세요.")
            return

        answer = QMessageBox.question(
            self,
            "삭제 확인",
            f"체크한 {len(product_ids)}개 상품을 삭제하시겠습니까?",
        )
        if answer == QMessageBox.StandardButton.Yes:
            self.database.delete_products(product_ids)
            self.clear_inputs()
            self.load_products()

    def clear_inputs(self):
        self.name_input.clear()
        self.price_input.setValue(0)
        self.product_table.clearSelection()

    def export_to_excel(self):
        if not self.products:
            QMessageBox.information(self, "저장 안내", "저장할 상품 데이터가 없습니다.")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "엑셀 파일 저장",
            "products.xlsx",
            "Excel 파일 (*.xlsx)",
        )
        if not file_path:
            return

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Products"
        worksheet.append(["productID", "productName", "productPrice"])
        for product in self.products:
            worksheet.append(
                [
                    product["productID"],
                    product["productName"],
                    product["productPrice"],
                ]
            )
        worksheet.column_dimensions["A"].width = 14
        worksheet.column_dimensions["B"].width = 24
        worksheet.column_dimensions["C"].width = 16

        try:
            workbook.save(file_path)
        except OSError as error:
            QMessageBox.critical(self, "저장 오류", f"엑셀 파일을 저장할 수 없습니다.\n{error}")
            return
        QMessageBox.information(self, "저장 완료", f"엑셀 파일을 저장했습니다.\n{file_path}")


if __name__ == "__main__":
    app = QApplication([])
    window = ProductWindow()
    window.show()
    app.exec()
