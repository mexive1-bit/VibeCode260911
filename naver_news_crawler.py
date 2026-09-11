from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
import requests
from bs4 import BeautifulSoup
from PyQt6.QtCore import QObject, QThread, pyqtSignal
from PyQt6.QtWidgets import (
    QApplication,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QSpinBox,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)


NAVER_SEARCH_URL = "https://search.naver.com/search.naver"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/131.0.0.0 Safari/537.36"
    )
}


@dataclass
class NewsArticle:
    title: str
    url: str
    summary: str
    press: str = ""
    published_at: str = ""
    naver_url: str = ""
    content: str = ""


def get_soup(url, params=None):
    response = requests.get(
        url,
        params=params,
        headers=HEADERS,
        timeout=15,
    )
    response.raise_for_status()
    response.encoding = response.apparent_encoding
    return BeautifulSoup(response.text, "html.parser")


def clean_text(element):
    if element is None:
        return ""
    return " ".join(element.get_text(" ", strip=True).split())


def search_news(keyword, display=10):
    soup = get_soup(
        NAVER_SEARCH_URL,
        params={
            "where": "news",
            "query": keyword,
            "start": 1,
            "display": display,
        },
    )

    articles = []
    seen_urls = set()

    # 제목과 요약은 제공된 HTML의 data-nlog-area 속성으로 연결한다.
    title_links = soup.select('a[data-nlog-area="nws_all.h.tit"]')
    if not title_links:
        # 네이버가 같은 결과를 클래스 기반 HTML로 내려주는 경우의 fallback이다.
        title_links = [
            title.find_parent("a", href=True)
            for title in soup.select(".sds-comps-text-type-headline1")
        ]

    for title_link in title_links:
        if title_link is None:
            continue

        url = title_link.get("href", "").strip()
        if not url.startswith(("http://", "https://")):
            continue

        title_block = title_link.parent
        summary_link = title_block.select_one('a[data-nlog-area="nws_all.h.body"]')
        if summary_link is None:
            summary_link = title_block.select_one(
                "a:has(.sds-comps-text-type-body1)"
            )
        article_block = title_link.find_parent(
            lambda tag: (
                tag.name == "div"
                and tag.select_one('a[data-nlog-area="nws_all.h.prof"]') is not None
            )
        )
        if article_block is None:
            article_block = title_block

        press_link = article_block.select_one('a[data-nlog-area="nws_all.h.prof"]')
        naver_link = article_block.select_one('a[data-nlog-area="nws_all.h.nav"]')
        profile_texts = article_block.select(
            ".sds-comps-profile-info-subtexts span.sds-comps-text"
        )
        title_text = title_link.select_one(
            ".sds-comps-text-type-headline1"
        ) or title_link
        summary_text = summary_link.select_one(
            ".sds-comps-text-type-body1"
        ) if summary_link else None
        press_text = (
            press_link.select_one(".sds-comps-profile-info-title-text")
            if press_link
            else None
        )

        add_article(
            articles,
            seen_urls,
            title=clean_text(title_text),
            url=url,
            summary=clean_text(summary_text or summary_link),
            press=clean_text(press_text or press_link),
            published_at=clean_text(profile_texts[0]) if profile_texts else "",
            naver_url=naver_link.get("href", "").strip() if naver_link else "",
        )

    return articles


def add_article(
    articles,
    seen_urls,
    title,
    url,
    summary,
    press="",
    published_at="",
    naver_url="",
):
    if not url or url in seen_urls:
        return

    seen_urls.add(url)
    articles.append(
        NewsArticle(
            title=title,
            url=url,
            summary=summary,
            press=press,
            published_at=published_at,
            naver_url=naver_url,
        )
    )


def get_article_content(url):
    soup = get_soup(url)

    # 언론사별로 본문 클래스가 달라 자주 사용되는 선택자를 순서대로 확인한다.
    selectors = [
        "article",
        ".article_body",
        ".article_body_contents",
        ".news_end",
        "#articleBodyContents",
        ".view_content",
    ]
    for selector in selectors:
        content = clean_text(soup.select_one(selector))
        if len(content) >= 80:
            return content

    return "본문을 찾지 못했습니다. 원문 사이트의 선택자를 확인하세요."


def crawl_news(keyword, display=10):
    articles = search_news(keyword, display)
    for article in articles:
        try:
            article.content = get_article_content(article.url)
        except requests.RequestException as error:
            article.content = f"원문 요청 실패: {error}"
    return articles


def save_to_excel(articles, file_path="naverResult.xlsx"):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "네이버 뉴스"

    headers = [
        "번호",
        "제목",
        "언론사",
        "작성 시간",
        "요약",
        "원문 링크",
        "네이버뉴스 링크",
        "본문",
    ]
    worksheet.append(headers)

    for index, article in enumerate(articles, start=1):
        worksheet.append(
            [
                index,
                article.title,
                article.press,
                article.published_at,
                article.summary,
                article.url,
                article.naver_url,
                article.content,
            ]
        )

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 8
    worksheet.column_dimensions["B"].width = 45
    worksheet.column_dimensions["C"].width = 15
    worksheet.column_dimensions["D"].width = 15
    worksheet.column_dimensions["E"].width = 60
    worksheet.column_dimensions["F"].width = 45
    worksheet.column_dimensions["G"].width = 45
    worksheet.column_dimensions["H"].width = 100

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].alignment = Alignment(horizontal="center", vertical="top")

    output_path = Path(file_path)
    workbook.save(output_path)
    return output_path


class CrawlWorker(QObject):
    completed = pyqtSignal(list, str)
    failed = pyqtSignal(str)

    def __init__(self, keyword, display):
        super().__init__()
        self.keyword = keyword
        self.display = display

    def run(self):
        try:
            articles = crawl_news(self.keyword, self.display)
            output_path = save_to_excel(articles)
            self.completed.emit(articles, str(output_path.resolve()))
        except Exception as error:
            self.failed.emit(str(error))


class NewsCrawlerWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.thread = None
        self.worker = None
        self.setup_ui()

    def setup_ui(self):
        self.setWindowTitle("네이버 뉴스 크롤러")
        self.resize(1200, 720)

        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)
        main_layout.setContentsMargins(20, 20, 20, 20)
        main_layout.setSpacing(12)

        title = QLabel("네이버 뉴스 크롤러")
        title.setObjectName("title")
        description = QLabel("검색어로 뉴스를 수집하고 naverResult.xlsx로 저장합니다.")
        description.setObjectName("description")
        main_layout.addWidget(title)
        main_layout.addWidget(description)

        form_layout = QFormLayout()
        self.keyword_input = QLineEdit("반도체")
        self.keyword_input.setPlaceholderText("검색어를 입력하세요")
        self.display_input = QSpinBox()
        self.display_input.setRange(1, 100)
        self.display_input.setValue(10)
        self.display_input.setSuffix(" 건")
        form_layout.addRow("검색어", self.keyword_input)
        form_layout.addRow("수집 개수", self.display_input)
        main_layout.addLayout(form_layout)

        button_layout = QHBoxLayout()
        self.crawl_button = QPushButton("뉴스 크롤링 시작")
        self.crawl_button.clicked.connect(self.start_crawling)
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setVisible(False)
        button_layout.addWidget(self.crawl_button)
        button_layout.addWidget(self.progress_bar)
        main_layout.addLayout(button_layout)

        self.status_label = QLabel("검색어와 수집 개수를 입력한 뒤 시작하세요.")
        main_layout.addWidget(self.status_label)

        self.result_table = QTableWidget(0, 5)
        self.result_table.setHorizontalHeaderLabels(
            ["번호", "제목", "언론사", "작성 시간", "요약"]
        )
        self.result_table.setAlternatingRowColors(True)
        self.result_table.setWordWrap(True)
        self.result_table.setColumnWidth(0, 60)
        self.result_table.setColumnWidth(1, 360)
        self.result_table.setColumnWidth(2, 130)
        self.result_table.setColumnWidth(3, 120)
        self.result_table.setColumnWidth(4, 520)
        self.result_table.verticalHeader().setDefaultSectionSize(64)
        main_layout.addWidget(self.result_table)

        self.setStyleSheet(
            """
            QMainWindow, QWidget {
                background-color: #f4f6f8;
                color: #263238;
                font-family: 'Malgun Gothic';
                font-size: 10pt;
            }
            QLabel#title {
                color: #17365d;
                font-size: 20pt;
                font-weight: 700;
            }
            QLabel#description, QLabel {
                color: #52606d;
            }
            QLineEdit, QSpinBox {
                background-color: white;
                border: 1px solid #b8c4ce;
                border-radius: 4px;
                padding: 7px;
            }
            QPushButton {
                background-color: #1f5f8b;
                border: 1px solid #1f5f8b;
                border-radius: 4px;
                color: white;
                font-weight: 700;
                padding: 8px 16px;
            }
            QPushButton:hover {
                background-color: #17496d;
            }
            QPushButton:disabled {
                background-color: #9aa7b2;
                border-color: #9aa7b2;
            }
            QTableWidget {
                background-color: white;
                alternate-background-color: #f7f9fb;
                border: 1px solid #d4dce3;
                gridline-color: #e1e7ec;
            }
            QHeaderView::section {
                background-color: #e7edf3;
                color: #34495e;
                font-weight: 700;
                padding: 8px;
                border: 0;
            }
            """
        )

    def start_crawling(self):
        keyword = self.keyword_input.text().strip()
        if not keyword:
            QMessageBox.warning(self, "입력 확인", "검색어를 입력하세요.")
            return

        self.crawl_button.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.status_label.setText(f"'{keyword}' 뉴스를 수집하는 중입니다...")
        self.result_table.setRowCount(0)

        self.thread = QThread()
        self.worker = CrawlWorker(keyword, self.display_input.value())
        self.worker.moveToThread(self.thread)
        self.thread.started.connect(self.worker.run)
        self.worker.completed.connect(self.show_results)
        self.worker.failed.connect(self.show_error)
        self.worker.completed.connect(self.thread.quit)
        self.worker.failed.connect(self.thread.quit)
        self.thread.finished.connect(self.crawling_finished)
        self.thread.start()

    def show_results(self, articles, output_path):
        self.result_table.setRowCount(len(articles))
        for row, article in enumerate(articles):
            values = [
                str(row + 1),
                article.title,
                article.press,
                article.published_at,
                article.summary,
            ]
            for column, value in enumerate(values):
                self.result_table.setItem(row, column, QTableWidgetItem(value))

        self.status_label.setText(
            f"{len(articles)}건 수집 완료. 저장 파일: {output_path}"
        )

    def show_error(self, message):
        self.status_label.setText("크롤링에 실패했습니다.")
        QMessageBox.critical(self, "크롤링 오류", message)

    def crawling_finished(self):
        self.progress_bar.setVisible(False)
        self.crawl_button.setEnabled(True)
        self.worker.deleteLater()
        self.thread.deleteLater()
        self.worker = None
        self.thread = None


def main():
    app = QApplication([])
    window = NewsCrawlerWindow()
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
